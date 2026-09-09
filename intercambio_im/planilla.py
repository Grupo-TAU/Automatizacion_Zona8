"""
Lectura y escritura de las planillas del intercambio con la IM.

Es la parte especifica del intercambio: el envio, la devolucion y el reporte, con
su hoja _meta y su proteccion de celdas. La lectura generica de planillas, que no
sabe nada de permisos, esta en tabla.py.
"""

from __future__ import annotations

from . import nucleo
from .nucleo import ErrorIntercambio, exigir_openpyxl

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
except ImportError:                   # pragma: no cover
    # No se rompe al importar: el plugin tiene que poder cargar igual y fallar
    # recien cuando alguien corra un algoritmo, con el mensaje de nucleo.
    Workbook = load_workbook = None
    Alignment = Font = PatternFill = Protection = get_column_letter = None


# Orden fijo de las filas de _meta. La lectura no depende del orden, pero un
# archivo devuelto es mas facil de auditar si siempre se ve igual.
CLAVES_META = (
    "lote_id", "fecha_export", "capa", "campo_clave",
    "proveedor", "fuente", "filas", "hash_claves",
)

# Las claves enviadas van en una columna aparte de _meta y no en una celda con
# JSON: el hash no permite reconstruir que filas se mandaron, y sin esa lista no
# se puede reportar 'no_devuelta'. Una celda tiene tope de 32767 caracteres; una
# columna, un millon de filas.
_COL_CLAVES_ENVIADAS = 4          # columna D
_TITULO_CLAVES = "claves_enviadas"

_RELLENO_ENCABEZADO = "FFD9D9D9"
_RELLENO_EDITABLE = "FFFFF2CC"


def _encabezar(hoja, titulos, fila=1) -> None:
    for columna, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(row=fila, column=columna, value=titulo)
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor=_RELLENO_ENCABEZADO)
        celda.alignment = Alignment(horizontal="center")


def _ajustar_anchos(hoja, anchos) -> None:
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho


# -- Escritura del envio (algoritmo A) ----------------------------------------
def escribir_envio(ruta, filas, meta: dict) -> None:
    """
    filas: [{'clave':…, 'N_Problema':…, 'Permisos_UCCRIU':…}] en el orden de envio.
    meta : los valores de CLAVES_META ya resueltos por el algoritmo.

    La hoja Datos queda protegida sin contrasena y con la columna del permiso
    desbloqueada. Es una guia visual para la IM, no seguridad: la proteccion se
    saca desde Excel en dos clics. La garantia real de que solo se toca esa
    columna es la lista blanca del algoritmo B, que no lee ninguna otra.
    """
    exigir_openpyxl()
    es_fid = meta.get("campo_clave") == nucleo.CLAVE_FID

    libro = Workbook()
    datos = libro.active
    datos.title = nucleo.HOJA_DATOS
    _encabezar(datos, nucleo.COLUMNAS_DATOS)

    for numero, fila in enumerate(filas, start=2):
        clave = fila[nucleo.COLUMNA_CLAVE]
        celda_clave = datos.cell(row=numero, column=1)
        if es_fid:
            celda_clave.value = int(clave)
        else:
            celda_clave.value = str(clave)
            celda_clave.number_format = "@"

        # Texto explicito en las dos columnas de texto: sin esto Excel se come
        # los ceros a la izquierda y convierte en fecha cualquier cosa que se le
        # parezca. En esta capa los permisos son texto libre ('655579-1',
        # '648493-1 / BAJA', 'BAJA'), asi que el riesgo es real.
        contexto = datos.cell(row=numero, column=2,
                              value=nucleo.texto_limpio(fila.get(nucleo.CAMPO_CONTEXTO)))
        contexto.number_format = "@"

        permiso = datos.cell(row=numero, column=3,
                             value=nucleo.texto_limpio(fila.get(nucleo.CAMPO_PERMISO)))
        permiso.number_format = "@"
        permiso.protection = Protection(locked=False)
        permiso.fill = PatternFill("solid", fgColor=_RELLENO_EDITABLE)

    _ajustar_anchos(datos, (12, 16, 26))
    datos.freeze_panes = "A2"
    # Sin password: asignarla, aunque sea None, hace que openpyxl la hashee.
    datos.protection.sheet = True

    meta_hoja = libro.create_sheet(nucleo.HOJA_META)
    _encabezar(meta_hoja, ("clave", "valor"))
    for numero, clave in enumerate(CLAVES_META, start=2):
        meta_hoja.cell(row=numero, column=1, value=clave)
        celda = meta_hoja.cell(row=numero, column=2, value=str(meta.get(clave, "")))
        celda.number_format = "@"

    celda = meta_hoja.cell(row=1, column=_COL_CLAVES_ENVIADAS, value=_TITULO_CLAVES)
    celda.font = Font(bold=True)
    celda.fill = PatternFill("solid", fgColor=_RELLENO_ENCABEZADO)
    for numero, fila in enumerate(filas, start=2):
        enviada = meta_hoja.cell(row=numero, column=_COL_CLAVES_ENVIADAS,
                                 value=str(fila[nucleo.COLUMNA_CLAVE]))
        enviada.number_format = "@"

    _ajustar_anchos(meta_hoja, (16, 62, 4, 18))
    meta_hoja.sheet_state = "hidden"

    libro.save(str(ruta))


# -- Lectura de la devolucion (algoritmo B) -----------------------------------
def leer_devolucion(ruta) -> tuple:
    """
    Devuelve (filas, meta, claves_enviadas) con todo CRUDO, sin normalizar: la
    normalizacion es responsabilidad de nucleo, para que sea la misma funcion de
    los dos lados.

    filas: [{'fila_excel': int, 'clave':…, 'N_Problema':…, 'Permisos_UCCRIU':…}]
    meta : {} si la IM borro la hoja _meta.
    """
    exigir_openpyxl()
    try:
        # data_only=True devuelve el ultimo valor calculado de las formulas. Si
        # la IM dejo formulas y guardo con algo que no las evalua, esas celdas
        # llegan como None y se reportan como vacias, nunca como un valor falso.
        libro = load_workbook(str(ruta), data_only=True)
    except Exception as exc:
        raise ErrorIntercambio(f"No se pudo abrir el XLSX devuelto: {exc}") from exc

    if nucleo.HOJA_DATOS not in libro.sheetnames:
        raise ErrorIntercambio(
            f"El archivo no tiene la hoja '{nucleo.HOJA_DATOS}'. "
            f"Hojas encontradas: {', '.join(libro.sheetnames) or '(ninguna)'}"
        )

    datos = libro[nucleo.HOJA_DATOS]
    columnas = _mapear_columnas(datos)
    filas = []
    for numero in range(2, datos.max_row + 1):
        fila = {
            nombre: datos.cell(row=numero, column=indice).value
            for nombre, indice in columnas.items()
        }
        # Una fila entera en blanco es el relleno que deja Excel al final, no un
        # dato que la IM haya borrado.
        if all(nucleo.esta_vacio(v) for v in fila.values()):
            continue
        fila["fila_excel"] = numero
        filas.append(fila)

    meta, claves = _leer_meta(libro)
    return filas, meta, claves


def _mapear_columnas(hoja) -> dict:
    """{nombre_esperado: indice_de_columna}, tolerando orden y mayusculas."""
    presentes = {}
    for indice in range(1, hoja.max_column + 1):
        titulo = nucleo.normalizar(hoja.cell(row=1, column=indice).value)
        if titulo and titulo not in presentes:
            presentes[titulo] = indice

    columnas = {}
    faltantes = []
    for esperada in nucleo.COLUMNAS_DATOS:
        indice = presentes.get(nucleo.normalizar(esperada))
        if indice is None:
            faltantes.append(esperada)
        else:
            columnas[esperada] = indice

    if faltantes:
        vistas = ", ".join(
            str(hoja.cell(row=1, column=i).value)
            for i in range(1, hoja.max_column + 1)
        )
        raise ErrorIntercambio(
            f"A la hoja '{nucleo.HOJA_DATOS}' le faltan columnas: "
            f"{', '.join(faltantes)}.\nEncabezados encontrados: {vistas or '(ninguno)'}"
        )
    return columnas


def _leer_meta(libro) -> tuple:
    if nucleo.HOJA_META not in libro.sheetnames:
        return {}, []

    hoja = libro[nucleo.HOJA_META]
    meta = {}
    for numero in range(2, hoja.max_row + 1):
        clave = nucleo.texto_limpio(hoja.cell(row=numero, column=1).value)
        if clave:
            meta[clave] = hoja.cell(row=numero, column=2).value

    claves = []
    if nucleo.normalizar(hoja.cell(row=1, column=_COL_CLAVES_ENVIADAS).value) == _TITULO_CLAVES:
        for numero in range(2, hoja.max_row + 1):
            valor = hoja.cell(row=numero, column=_COL_CLAVES_ENVIADAS).value
            if not nucleo.esta_vacio(valor):
                claves.append(valor)
    return meta, claves


# -- Reporte (algoritmo B) -----------------------------------------------------
COLUMNAS_REPORTE = ("clave", nucleo.CAMPO_CONTEXTO, "valor_anterior",
                    "valor_nuevo", "accion")


def escribir_reporte(ruta, evaluaciones, conteos: dict, encabezado: dict) -> None:
    """Una fila por evaluacion, mas una hoja con el resumen y el contexto de la corrida."""
    exigir_openpyxl()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    _encabezar(hoja, COLUMNAS_REPORTE)

    for numero, evaluacion in enumerate(evaluaciones, start=2):
        hoja.cell(row=numero, column=1, value=str(evaluacion.clave)).number_format = "@"
        hoja.cell(row=numero, column=2, value=evaluacion.n_problema).number_format = "@"
        hoja.cell(row=numero, column=3, value=evaluacion.valor_anterior).number_format = "@"
        hoja.cell(row=numero, column=4, value=evaluacion.valor_nuevo).number_format = "@"
        hoja.cell(row=numero, column=5, value=evaluacion.accion)

    _ajustar_anchos(hoja, (12, 26, 22, 22, 16))
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = f"A1:E{max(len(evaluaciones) + 1, 1)}"

    resumen = libro.create_sheet("Resumen")
    _encabezar(resumen, ("dato", "valor"))
    numero = 2
    for clave, valor in encabezado.items():
        resumen.cell(row=numero, column=1, value=clave)
        resumen.cell(row=numero, column=2, value=str(valor)).number_format = "@"
        numero += 1
    numero += 1
    resumen.cell(row=numero, column=1, value="accion").font = Font(bold=True)
    resumen.cell(row=numero, column=2, value="filas").font = Font(bold=True)
    numero += 1
    for accion in nucleo.ACCIONES:
        resumen.cell(row=numero, column=1, value=accion)
        resumen.cell(row=numero, column=2, value=conteos.get(accion, 0))
        numero += 1

    _ajustar_anchos(resumen, (22, 62))
    libro.save(str(ruta))
